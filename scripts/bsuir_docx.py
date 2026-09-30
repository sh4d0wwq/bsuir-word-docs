"""BSUIR STP 01-2017 .docx builder.
  init CONFIG.json OUT.docx   - styled document + title page (+ abstract, contents)
  add  DOC.docx CHUNK.md      - append markup chunk (see SKILL.md)
  fix  DOC.docx [--pdf]       - normalize MCP-added content, lint, update fields and check layout via Word
  build CONFIG.json OUT.docx CHUNK.md... [--pdf] - init + add all chunks + fix in one call
  preview DOC.docx|PDF OUTDIR [1,2,8] - contact sheet(s) of all pages, or listed pages at full size
"""
import contextlib, io, json, os, re, struct, subprocess, sys, tempfile, zlib, copy
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

FONT, LINE, TEXT_W, TEXT_H = "Times New Roman", Pt(18), Cm(16.5), Cm(21)
HEADINGS = {"Heading 1", "Heading 2", "Heading 3", "StructHeading", "PlainHeading", "AppendixHeading"}


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def _font(style, size=None, bold=None, caps=None, font=FONT):
    rpr = style.element.get_or_add_rPr()
    for tag in ("w:rFonts", "w:color"):
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
    rpr.insert(0, _el("w:rFonts", ascii=font, hAnsi=font, cs=font, eastAsia=font))
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    if size: style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if caps is not None: style.font.all_caps = caps


def _style(doc, name, base="Normal", size=None, bold=None, caps=None, align=None, fi=None, li=None,
           sb=None, sa=None, ls=None, kwn=None, pbb=None, outline=None, tabs=(), font=FONT, nohyph=False):
    try:
        s = doc.styles[name]
    except KeyError:
        s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if name != "Normal":
        s.base_style = doc.styles[base]
    _font(s, size, bold, caps, font)
    f = s.paragraph_format
    for attr, val in (("alignment", align), ("first_line_indent", fi), ("left_indent", li), ("space_before", sb),
                      ("space_after", sa), ("line_spacing", ls), ("keep_with_next", kwn),
                      ("page_break_before", pbb)):
        if val is not None:
            setattr(f, attr, val)
    for pos, al, lead in tabs:
        f.tab_stops.add_tab_stop(pos, al, lead)
    ppr = s.element.get_or_add_pPr()
    if outline is not None:
        for o in ppr.findall(qn("w:outlineLvl")):
            ppr.remove(o)
        ppr.append(_el("w:outlineLvl", val=outline))
    if nohyph:
        ppr.append(_el("w:suppressAutoHyphens"))
    return s


def setup_styles(doc, new_page_sections=True):
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rpr.append(_el("w:lang", val="ru-RU", eastAsia="ru-RU", bidi="ar-SA"))
    doc.settings.element.append(_el("w:autoHyphenation"))
    _style(doc, "Normal", size=14, align=AL.JUSTIFY, fi=Cm(1.25), li=0, sb=0, sa=0, ls=1.1)
    doc.styles["Normal"].paragraph_format.widow_control = True
    _style(doc, "Heading 1", size=14, bold=True, caps=True, align=AL.LEFT, fi=0, li=Cm(1.25),
           sb=0 if new_page_sections else LINE, sa=LINE, kwn=True, pbb=new_page_sections, nohyph=True)
    for h in ("Heading 2", "Heading 3"):
        _style(doc, h, size=14, bold=True, caps=False, align=AL.LEFT, fi=0, li=Cm(1.25), sb=LINE, sa=LINE,
               kwn=True, nohyph=True)
    _style(doc, "StructHeading", "Heading 1", align=AL.CENTER, li=0, sb=0, pbb=True, outline=0)
    _style(doc, "PlainHeading", "Normal", size=14, bold=True, caps=True, align=AL.CENTER, fi=0, sa=LINE,
           kwn=True, nohyph=True)
    _style(doc, "AppendixHeading", "Heading 1", caps=False, align=AL.CENTER, li=0, sb=0, pbb=True, outline=0)
    _style(doc, "Figure", align=AL.CENTER, fi=0, sb=LINE, sa=LINE, ls=1.0, kwn=True)
    _style(doc, "FigureCaption", align=AL.CENTER, fi=0, sa=LINE, nohyph=True)
    _style(doc, "TableCaption", align=AL.LEFT, fi=0, sb=LINE, kwn=True, nohyph=True)
    _style(doc, "TableText", size=14, align=AL.LEFT, fi=0, ls=1.1)
    _style(doc, "Formula", fi=0, sb=LINE, sa=LINE,
           tabs=((Cm(8.25), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.SPACES),
                 (TEXT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)))
    _style(doc, "Where", fi=Cm(-1), li=Cm(1), align=AL.LEFT)
    _style(doc, "Sublist", fi=Cm(2.5))
    _style(doc, "Code", size=12, align=AL.LEFT, fi=0, ls=1.0, font="Courier New", nohyph=True)
    for extra in ("Header", "Footer"):
        try:
            _font(doc.styles[extra], 14)
        except KeyError:
            pass


INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")


def runs(p, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*"):
            p.add_run(part[1:-1]).italic = True
        elif part.startswith("_{"):
            p.add_run(part[2:-1]).font.subscript = True
        elif part.startswith("^{"):
            p.add_run(part[2:-1]).font.superscript = True
        else:
            p.add_run(part)
    return p


def para(doc, text="", style="Normal", raw=False):
    p = doc.add_paragraph(style=style)
    if raw:
        p.add_run(text)
    else:
        runs(p, text)
    return p


def save(doc, path):
    try:
        doc.save(path)
    except PermissionError:
        sys.exit(f"error: {path} is locked (open in Word?). Close it or build to another path.")


def field(p, instr, placeholder=""):
    r = p.add_run()
    r._r.append(_el("w:fldChar", fldCharType="begin"))
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r._r.append(it)
    r._r.append(_el("w:fldChar", fldCharType="separate"))
    p.add_run(placeholder)
    p.add_run()._r.append(_el("w:fldChar", fldCharType="end"))


def _plain(p, align=AL.CENTER, li=None):
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.alignment = align
    if li is not None:
        p.paragraph_format.left_indent = li
    return p


TITLES = {
    "lab": "**ОТЧЕТ**\nпо лабораторной работе № {number}\nпо дисциплине «{discipline}»\nна тему\n«{topic}»",
    "course": "**ПОЯСНИТЕЛЬНАЯ ЗАПИСКА**\nк курсовому проекту\nпо дисциплине «{discipline}»\nна тему\n**{TOPIC}**",
    "diploma": "**ПОЯСНИТЕЛЬНАЯ ЗАПИСКА**\nк дипломному проекту\nна тему\n**{TOPIC}**",
}


def _page_numbers(section):
    section.footer.is_linked_to_previous = False
    field(_plain(section.footer.paragraphs[0], AL.RIGHT), "PAGE", "1")


def cmd_init(cfg_path, out):
    c = json.load(open(cfg_path, encoding="utf-8-sig"))
    kind = c.get("type", "lab")
    c.setdefault("number", ""); c.setdefault("discipline", ""); c.setdefault("topic", "")
    c["TOPIC"] = c["topic"].upper()
    doc = Document()
    setup_styles(doc, c.get("section_new_page", True))
    s0 = doc.sections[0]
    s0.page_width, s0.page_height = Mm(210), Mm(297)
    s0.top_margin, s0.bottom_margin, s0.left_margin, s0.right_margin = Mm(20), Mm(27), Mm(30), Mm(15)
    s0.footer_distance = Mm(17)

    # title page: blocks spread vertically (vAlign=both)
    _plain(para(doc, c.get("header", "Министерство образования Республики Беларусь\n\nУчреждение образования\n"
                                "БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ\nИНФОРМАТИКИ И РАДИОЭЛЕКТРОНИКИ")))
    fac = "\n".join(x for x in (c.get("faculty"), c.get("department")) if x)
    if fac:
        _plain(para(doc, fac), AL.LEFT if kind == "diploma" else AL.CENTER)
    if c.get("head"):
        _plain(para(doc, f"К защите допустить:\nЗаведующий кафедрой {c.get('dept_short', '')}\n"
                         f"____________{c['head']}"), AL.LEFT, Cm(9))
    _plain(para(doc, c.get("title_block") or TITLES[kind].format(**c)))
    if c.get("code"):
        _plain(para(doc, c["code"]))
    signers = c.get("signers")
    if not signers and c.get("student"):
        g = f" гр. {c['group']}" if c.get("group") else ""
        signers = [[f"Выполнил: студент{g}" if kind == "lab" else f"Студент{g}", c["student"]]]
        if c.get("teacher"):
            signers.append(["Проверил:" if kind == "lab" else "Руководитель", c["teacher"]])
    if signers:
        p = _plain(para(doc, "\n".join(f"{a}\t{b}" for a, b in signers)), AL.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(c.get("sign_tab_cm", 10)))
    _plain(para(doc, f"{c.get('city', 'Минск')} {c.get('year', '')}".strip()))
    s0._sectPr.find(qn("w:cols")).addnext(_el("w:vAlign", val="both"))

    new = doc.add_section(WD_SECTION.NEW_PAGE)
    va = new._sectPr.find(qn("w:vAlign"))
    if va is not None:
        new._sectPr.remove(va)
    toc = c.get("toc", kind != "lab")
    if toc:
        _page_numbers(new)
        para(doc, "Содержание", "PlainHeading")
        field(_plain(doc.add_paragraph(), AL.LEFT), 'TOC \\o "1-3" \\h \\z \\u', "Запустите fix для обновления")
    if c.get("abstract") or c.get("assignment_pages"):
        if c.get("abstract"):
            h = para(doc, "Реферат", "PlainHeading")
            h.paragraph_format.page_break_before = toc
            for t in ([c["abstract_header"]] if c.get("abstract_header") else []) + c["abstract"].split("\n"):
                para(doc, t)
        for i in range(c.get("assignment_pages", 0)):
            p = _plain(para(doc, f"[Бланк задания, лист {i + 1}]"))
            p.paragraph_format.page_break_before = True
        if not toc:
            new = doc.add_section(WD_SECTION.NEW_PAGE)
    if not toc:
        _page_numbers(new)
    doc.core_properties.author = c.get("student", "")
    doc.core_properties.title = c.get("topic", "")
    save(doc, out)
    print("ok", out)


RE_IMG = re.compile(r"^!\[.*?\]\((.+?)\)(?:\{w=([\d.]+)\})?$")
RE_FORM = re.compile(r"^\$\$(.+?)\$\$\s*(\(.+?\))?$")


def _dummy_png():
    """Gray framed PNG used instead of a real figure."""
    path = os.path.join(tempfile.gettempdir(), "bsuir_figure_placeholder.png")
    if os.path.isfile(path) and os.path.getsize(path) > 0:
        return path
    w, h = 900, 560
    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    raw = bytearray()
    for y in range(h):
        raw.append(0)
        for x in range(w):
            edge = x < 4 or y < 4 or x >= w - 4 or y >= h - 4
            raw += b"\x66\x66\x66" if edge else b"\xe6\xe6\xe6"
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(bytes(raw), 9)) + chunk(b"IEND", b"")
    open(path, "wb").write(png)
    return path


def add_table(doc, rows):
    rows = [r for r in rows if not re.fullmatch(r"[\s|:\-]+", r)]
    cells = [[x.strip() for x in r.strip().strip("|").split("|")] for r in rows]
    n = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=n)
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j in range(n):
            p = t.cell(i, j).paragraphs[0]
            p.style = "TableText"
            runs(p, r[j] if j < len(r) and r[j] else "–")


def load_code(path, base):
    path = path.strip() if os.path.isabs(path.strip()) else os.path.join(base, path.strip())
    return open(path, encoding="utf-8").read().rstrip("\n").splitlines()


def _gap_before_code(doc):
    for el in reversed(doc.element.body):
        if el.tag == qn("w:tbl"):
            return
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc._body)
            if p.text.strip() and _sname(p) not in HEADINGS | {"Code", "Figure", "FigureCaption"}:
                doc.add_paragraph()
            return


def cmd_add(docp, md):
    doc = Document(docp)
    base = os.path.dirname(os.path.abspath(md))
    lines = open(md, encoding="utf-8-sig").read().splitlines()
    code, where, table, gap = False, False, [], False
    for line in lines + [""]:
        if code:
            if line.strip() == "```":
                code, gap = False, True
            else:
                para(doc, line, "Code", raw=True)
            continue
        s = line.rstrip()
        if s.lstrip().startswith("|"):
            table.append(s); continue
        if table:
            add_table(doc, table); table = []
        if s.startswith("```"):
            _gap_before_code(doc); code = True; continue
        if s.startswith("@code "):
            _gap_before_code(doc)
            for cl in load_code(s[6:], base):
                para(doc, cl, "Code", raw=True)
            gap = True; continue
        if not s.startswith("  "):
            where = False
        if not s.strip():
            continue
        m = RE_IMG.match(s); f = RE_FORM.match(s)
        if gap:
            if not (m or s.startswith("#") or s == "\\newpage" or re.match(r"^Таблица [\dА-Я]", s)):
                doc.add_paragraph()
            gap = False
        if m:
            width = min(float(m.group(2)) if m.group(2) else 14, 16.5)
            p = doc.add_paragraph(style="Figure")
            p.add_run().add_picture(_dummy_png(), width=Cm(width))
        elif f:
            para(doc, "\t" + f.group(1).strip() + ("\t" + f.group(2) if f.group(2) else ""), "Formula", raw=True)
        elif s.startswith("#@ "):
            letter, kind_, title = [x.strip() for x in s[3:].split("|")]
            para(doc, f"ПРИЛОЖЕНИЕ {letter}\n({kind_})\n{title}", "AppendixHeading")
        elif s.startswith("#! "):
            para(doc, s[3:], "StructHeading")
        elif s.startswith("#= "):
            para(doc, s[3:], "PlainHeading")
        elif s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            para(doc, s[lvl:].strip(), f"Heading {min(lvl, 3)}")
        elif s == "\\newpage":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif s.startswith("где "):
            para(doc, "где\t" + s[4:], "Where"); where = True
        elif where and s.startswith("  "):
            para(doc, "\t" + s.strip(), "Where")
        elif re.match(r"^Рисунок [\dА-Я]", s):
            para(doc, s, "FigureCaption")
        elif re.match(r"^Таблица [\dА-Я]", s):
            para(doc, s, "TableCaption")
        elif re.match(r"^\s{2,}(- |\d+\) |[а-я]\) )", s):
            st = s.strip()
            para(doc, "– " + st[2:] if st.startswith("- ") else st, "Sublist")
        elif s.startswith("- "):
            para(doc, "– " + s[2:])
        else:
            para(doc, s.strip())
    save(doc, docp)
    print("ok")


def _sname(p):
    try:
        return p.style.name
    except Exception:
        return "Normal"


LIST_KINDS = {"dash": ("bullet", "–"), "dec": ("decimal", "%1)"), "let": ("russianLower", "%1)")}
LIST_BASE = 900
RE_ITEM = re.compile(r"^(– |\d+\) |[а-я]\) )")


def _list_abstracts(numbering):
    """abstractNumId -> (kind, level); one single-level definition per kind and nesting level."""
    ids = {}
    for i, (kind, (fmt, txt)) in enumerate(LIST_KINDS.items()):
        for lvl in (0, 1):
            aid = LIST_BASE + i * 2 + lvl
            ids[aid] = (kind, lvl)
            if numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{aid}"]'):
                continue
            a = _el("w:abstractNum", abstractNumId=aid)
            a.append(_el("w:multiLevelType", val="singleLevel"))
            lv = _el("w:lvl", ilvl=0)
            for tag, val in (("start", 1), ("numFmt", fmt), ("suff", "space"), ("lvlText", txt), ("lvlJc", "left")):
                lv.append(_el("w:" + tag, val=val))
            ppr = _el("w:pPr"); ppr.append(_el("w:ind", left=0, firstLine=709 * (lvl + 1))); lv.append(ppr)
            rpr = _el("w:rPr"); rpr.append(_el("w:rFonts", ascii=FONT, hAnsi=FONT, cs=FONT, eastAsia=FONT)); lv.append(rpr)
            a.append(lv)
            first_num = numbering.find(qn("w:num"))
            if first_num is not None:
                first_num.addprevious(a)
            else:
                numbering.append(a)
    return ids


def fix_lists(doc):
    """Turn text list items («– », «1) », «а) » in Normal/Sublist) into Word numbering; restart per list."""
    numbering = doc.part.numbering_part.element
    abstracts = _list_abstracts(numbering)
    num_abs = {int(n.get(qn("w:numId"))): int(n.find(qn("w:abstractNumId")).get(qn("w:val")))
               for n in numbering.findall(qn("w:num"))}
    bullets, cur = {}, {}

    def new_num(aid, restart):
        n = numbering.add_num(aid)
        if restart:
            o = _el("w:lvlOverride", ilvl=0); o.append(_el("w:startOverride", val=1)); n.append(o)
        nid = int(n.get(qn("w:numId"))); num_abs[nid] = aid
        return nid

    for el in doc.element.body:
        if el.tag != qn("w:p"):
            cur.clear(); continue
        p = Paragraph(el, doc._body)
        numpr = el.pPr.numPr if el.pPr is not None else None
        if numpr is not None and numpr.numId is not None and num_abs.get(numpr.numId.val) in abstracts:
            kind, lvl = abstracts[num_abs[numpr.numId.val]]
            cur[lvl] = (kind, numpr.numId.val)
            if lvl == 0: cur.pop(1, None)
            continue
        sn, m = _sname(p), RE_ITEM.match(p.text)
        if not (m and sn in ("Normal", "Sublist")):
            cur.clear(); continue
        lvl = int(sn == "Sublist")
        kind = "dash" if m.group(1) == "– " else "dec" if m.group(1)[0].isdigit() else "let"
        aid = LIST_BASE + list(LIST_KINDS).index(kind) * 2 + lvl
        if kind == "dash":
            nid = bullets.get(aid) or bullets.setdefault(aid, new_num(aid, False))
        elif cur.get(lvl, (None,))[0] == kind:
            nid = cur[lvl][1]
        else:
            nid = new_num(aid, True)
        cur[lvl] = (kind, nid)
        if lvl == 0: cur.pop(1, None)
        cut = len(m.group(1))
        for r in p.runs:
            k = min(cut, len(r.text)); r.text = r.text[k:]; cut -= k
            if not cut: break
        np_ = el.get_or_add_pPr().get_or_add_numPr()
        np_.get_or_add_ilvl().val = 0
        np_.get_or_add_numId().val = nid


HEAD_ABS = 950
RE_HNUM = re.compile(r"^\s*(\d+(?:\.\d+)*)\s+")


def _heading_numbering(doc):
    """Multilevel list 1 / 1.1 / 1.1.1 linked to Heading 1-3; returns its numId."""
    numbering = doc.part.numbering_part.element
    for n in numbering.findall(qn("w:num")):
        if n.find(qn("w:abstractNumId")).get(qn("w:val")) == str(HEAD_ABS):
            return int(n.get(qn("w:numId")))
    a = _el("w:abstractNum", abstractNumId=HEAD_ABS)
    a.append(_el("w:multiLevelType", val="multilevel"))
    for i in range(3):
        lv = _el("w:lvl", ilvl=i)
        for tag, val in (("start", 1), ("numFmt", "decimal"), ("pStyle", doc.styles[f"Heading {i + 1}"].style_id),
                         ("suff", "space"), ("lvlText", ".".join(f"%{k + 1}" for k in range(i + 1))),
                         ("lvlJc", "left")):
            lv.append(_el("w:" + tag, val=val))
        ppr = _el("w:pPr"); ppr.append(_el("w:ind", left=709, firstLine=0)); lv.append(ppr)
        a.append(lv)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(a)
    else:
        numbering.append(a)
    return int(numbering.add_num(HEAD_ABS).get(qn("w:numId")))


def fix_headings(doc):
    """Auto-number Heading 1-3 (number typed in the text is removed and checked); appendix headings keep text numbers."""
    nid, warn = _heading_numbering(doc), []
    for i in range(3):
        np_ = doc.styles[f"Heading {i + 1}"].element.get_or_add_pPr().get_or_add_numPr()
        np_.get_or_add_ilvl().val = i
        np_.get_or_add_numId().val = nid
    for name in ("StructHeading", "AppendixHeading"):
        doc.styles[name].element.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
    cnt, in_app = [0, 0, 0], False
    for p in doc.paragraphs:
        sn = _sname(p)
        if sn in ("AppendixHeading", "StructHeading"):
            in_app = sn == "AppendixHeading"
        if sn not in ("Heading 1", "Heading 2", "Heading 3"):
            continue
        if in_app:
            p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(1.25), 0
            continue
        lvl = int(sn[-1]) - 1
        cnt[lvl] += 1
        cnt[lvl + 1:] = [0] * (2 - lvl)
        m = RE_HNUM.match(p.text)
        if not m:
            continue
        want = ".".join(map(str, cnt[:lvl + 1]))
        if m.group(1) != want:
            warn.append(f"номер заголовка {m.group(1)} заменен автономером {want}: {p.text[m.end():][:40]}")
        cut = m.end()
        for r in p.runs:
            k = min(cut, len(r.text)); r.text = r.text[k:]; cut -= k
            if not cut: break
    return warn


def cmd_fix(docp, pdf=False):
    doc = Document(docp)
    body = doc.element.body
    warn = fix_headings(doc)
    fix_lists(doc)
    # pictures added via MCP -> Figure style, clamp width
    for p in doc.paragraphs:
        if p._p.findall(".//" + qn("w:drawing")):
            if _sname(p) == "Normal":
                p.style = "Figure"
            for ext in p._p.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"):
                cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                k = min(1, TEXT_W / cx, TEXT_H / cy)
                if k < 1:
                    for e in p._p.iter():
                        if e.tag.endswith("}extent") or e.tag.endswith("}ext"):
                            if e.get("cx"):
                                e.set("cx", str(int(int(e.get("cx")) * k))); e.set("cy", str(int(int(e.get("cy")) * k)))
    # tables: exact text width, 16.5 cm between the 30 mm and 15 mm margins
    sec = doc.sections[-1]
    text_twips = int(round((sec.page_width - sec.left_margin - sec.right_margin) / 635))
    for t in doc.tables:
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl = t._tbl
        tblPr = tbl.tblPr
        for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
            for old in tblPr.findall(qn(tag)):
                tblPr.remove(old)
        tblPr.append(_el("w:tblW", w=text_twips, type="dxa"))
        tblPr.append(_el("w:tblInd", w=108, type="dxa"))
        tblPr.append(_el("w:tblLayout", type="fixed"))
        grid = tbl.find(qn("w:tblGrid"))
        cols = grid.findall(qn("w:gridCol")) if grid is not None else []
        n = len(cols) or len(t.columns)
        base, rem = divmod(text_twips, n)
        widths = [base + (1 if i < rem else 0) for i in range(n)]
        for col, width in zip(cols, widths):
            col.set(qn("w:w"), str(width))
        for row in t.rows:
            seen, ci = set(), 0
            for cell in row.cells:
                if id(cell._tc) in seen or ci >= n:
                    continue
                seen.add(id(cell._tc))
                tcPr = cell._tc.get_or_add_tcPr()
                for old in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(old)
                span = tcPr.find(qn("w:gridSpan"))
                k = int(span.get(qn("w:val"))) if span is not None else 1
                tcPr.append(_el("w:tcW", w=sum(widths[ci:ci + k]), type="dxa"))
                ci += k
        for i, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            if not trPr.findall(qn("w:cantSplit")):
                trPr.append(_el("w:cantSplit"))
            if i == 0 and not trPr.findall(qn("w:tblHeader")):
                trPr.append(_el("w:tblHeader"))
            for cell in row.cells:
                for p in cell.paragraphs:
                    if _sname(p) in ("Normal", "TableText"):
                        p.style = "TableText"
                        if i == 0:
                            p.alignment = AL.CENTER
        prev, nxt = t._tbl.getprevious(), t._tbl.getnext()
        if prev is None or prev.tag != qn("w:p") or _sname(Paragraph(prev, t._parent)) != "TableCaption":
            warn.append("таблица без заголовка «Таблица N – ...» над ней")
        if nxt is not None and nxt.tag == qn("w:p") and "".join(nxt.itertext()).strip():
            nxt.addprevious(OxmlElement("w:p"))
    # heading spacing, dashes, lint data
    prev = None
    text_all, figs, tabs, apps, cites, bib, in_bib, holes = [], [], [], [], [], [], False, []
    for p in doc.paragraphs:
        sn = _sname(p)
        if sn in ("Heading 2", "Heading 3") and prev in HEADINGS:
            p.paragraph_format.space_before = 0
        if sn != "Code":
            for r in p.runs:
                if " - " in r.text:
                    r.text = r.text.replace(" - ", " – ")
        tx = p.text.strip()
        holes.extend(re.findall(r"\[\[(.+?)\]\]", tx))
        if sn in HEADINGS:
            in_bib = "источник" in tx.lower()
            if tx.endswith(".") and sn != "AppendixHeading":
                warn.append(f"точка в конце заголовка: {tx[:50]}")
        if sn == "FigureCaption":
            figs.append(tx.split()[1] if len(tx.split()) > 1 else "?")
            if tx.endswith("."): warn.append(f"точка в конце подписи: {tx[:50]}")
        elif sn == "TableCaption":
            tabs.append(tx.split()[1] if len(tx.split()) > 1 else "?")
        elif sn == "AppendixHeading":
            apps.append(tx.split()[1] if len(tx.split()) > 1 else "?")
        elif in_bib and re.match(r"^\[\d+\]", tx):
            bib.append(int(re.match(r"^\[(\d+)\]", tx).group(1)))
        elif sn not in HEADINGS and sn != "Code":
            text_all.append(tx)
            for m in re.finditer(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]", tx):
                for part in re.split(r"\s*,\s*", m.group(1)):
                    a = re.split(r"\s*[–-]\s*", part)
                    cites.extend(range(int(a[0]), int(a[-1]) + 1))
        prev = sn
    joined = " ".join(text_all).lower().replace("\xa0", " ")
    for kind, nums, stem in (("рисунок", figs, "рисун"), ("таблица", tabs, "таблиц"), ("приложение", apps, "приложени")):
        for n in nums:
            if not re.search(stem + r"\w*\s+(\S+\s+)?" + re.escape(n.lower()), joined):
                warn.append(f"нет ссылки в тексте: {kind} {n}")
    first = list(dict.fromkeys(cites))
    if first != sorted(first):
        warn.append(f"источники цитируются не по порядку: {first[:15]}")
    for n in set(bib) - set(cites):
        warn.append(f"нет ссылки на источник [{n}]")
    for n in set(cites) - set(bib):
        if bib: warn.append(f"ссылка [{n}] без записи в списке")
    if holes:
        warn.append("плейсхолдеры, сообщите пользователю: " + "; ".join(dict.fromkeys(holes)))
    save(doc, docp)
    js = os.path.join(tempfile.gettempdir(), "bsuir_update.js")
    open(js, "w", encoding="utf-16").write(JS)
    args = ["cscript", "//nologo", js, os.path.abspath(docp)]
    if pdf:
        args.append(os.path.splitext(os.path.abspath(docp))[0] + ".pdf")
    try:
        raw = subprocess.run(args, capture_output=True, timeout=300).stdout
        try:
            out = raw.decode("utf-8").strip()
        except UnicodeDecodeError:
            out = raw.decode("cp866", "replace").strip()
    except Exception as e:
        out = f"Word COM недоступен ({e}); обновите поля в Word: Ctrl+A, F9"
    print(out)
    print("\n".join(dict.fromkeys(warn)) or "lint: ok")


JS = r"""var a=WScript.Arguments,w=new ActiveXObject("Word.Application"),d=null,f=0;
w.Visible=false;w.DisplayAlerts=0;
try{d=w.Documents.Open(a(0),false,false,false);
for(var i=1;i<=d.Paragraphs.Count;i++){var p=d.Paragraphs(i);
if(p.Style.NameLocal!="Formula"||p.Range.OMaths.Count>0)continue;
var r=p.Range,t=r.Text,x=t.indexOf("\t"),y=t.lastIndexOf("\t");if(x<0)continue;
var e=(y>x)?r.Start+y:r.End-1,m=d.Range(r.Start+x+1,e);d.OMaths.Add(m).OMaths(1).BuildUp();f++;}
for(var k=0;k<3;k++){var st=d.Styles(-20-k),pf=st.ParagraphFormat;pf.Alignment=0;pf.FirstLineIndent=0;
pf.LeftIndent=w.CentimetersToPoints(0.5*k);pf.RightIndent=0;st.Font.Name="Times New Roman";st.Font.Size=14;}
for(k=1;k<=d.TablesOfContents.Count;k++)d.TablesOfContents(k).Update();
d.Fields.Update();for(k=1;k<=d.TablesOfContents.Count;k++)d.TablesOfContents(k).Update();
var n=d.ComputeStatistics(2),L=[];
for(k=2;k<=n;k++){var s=d.GoTo(1,1,k).Start,e=(k<n)?d.GoTo(1,1,k+1).Start:d.Content.End,g=d.Range(s,e);
var c=g.Text.replace(/\s/g,"").length;if(c<150&&g.Paragraphs(1).OutlineLevel==10&&g.InlineShapes.Count==0&&g.Tables.Count==0)L.push("layout: стр. "+k+" почти пустая ("+c+" знаков)");}
for(k=1;k<=d.InlineShapes.Count;k++){var sp=d.InlineShapes(k).Range,pp=sp.Paragraphs(1),pg=sp.Information(3),nx=pp.Next(),pv=pp.Previous();
if(nx&&nx.Range.Information(3)!=pg)L.push("layout: стр. "+pg+": рисунок и подпись на разных страницах");
if(pv&&pv.Range.Information(3)<pg){var gap=(d.PageSetup.PageHeight-d.PageSetup.BottomMargin-pv.Range.Characters.Last.Information(6))/17.7;
if(gap>8)L.push("layout: стр. "+(pg-1)+": пусто ~"+Math.round(gap)+" строк внизу, рисунок перенесен на стр. "+pg+" (уменьшите w или поставьте следующий абзац перед рисунком)");}}
d.Save();if(a.length>1)d.ExportAsFixedFormat(a(1),17);
WScript.Echo("pages="+n+" formulas_converted="+f+(a.length>1?" pdf="+a(1):"")+(L.length?"\n"+L.join("\n"):""));}
catch(err){WScript.Echo("COM error: "+err.message);}
try{if(d)d.Close(0);}catch(e2){}w.Quit(0);"""

PDF_JS = r"""var a=WScript.Arguments,w=new ActiveXObject("Word.Application");w.Visible=false;w.DisplayAlerts=0;
try{var d=w.Documents.Open(a(0),false,true);d.ExportAsFixedFormat(a(1),17);d.Close(0);}
catch(e){WScript.Echo("COM error: "+e.message);}w.Quit(0);"""


def cmd_preview(docp, outdir, pages=""):
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    os.makedirs(outdir, exist_ok=True)
    pdf = os.path.abspath(docp)
    if not pdf.lower().endswith(".pdf"):
        js = os.path.join(tempfile.gettempdir(), "bsuir_pdf.js")
        open(js, "w", encoding="utf-16").write(PDF_JS)
        pdf = os.path.join(tempfile.gettempdir(), "bsuir_preview.pdf")
        subprocess.run(["cscript", "//nologo", js, os.path.abspath(docp), pdf], timeout=300)
    d = pymupdf.open(pdf)
    nums = [int(x) for x in pages.split(",") if x.strip()]
    if nums:
        for n in (x for x in nums if 1 <= x <= d.page_count):
            path = os.path.join(outdir, f"page_{n:02d}.png")
            d[n - 1].get_pixmap(matrix=pymupdf.Matrix(1.0, 1.0)).save(path)
            print(path)
    else:
        W, H, COLS, ROWS = 595, 842, 4, 3
        for s in range(0, d.page_count, COLS * ROWS):
            rows = min(ROWS, -(-(d.page_count - s) // COLS))
            sheet = pymupdf.open().new_page(width=W * COLS, height=H * rows)
            for i, n in enumerate(range(s, min(s + COLS * ROWS, d.page_count))):
                r = pymupdf.Rect((i % COLS) * W, (i // COLS) * H, (i % COLS + 1) * W, (i // COLS + 1) * H)
                sheet.show_pdf_page(r, d, n)
                sheet.draw_rect(r, color=(0.5, 0.5, 0.5), width=2)
            path = os.path.join(outdir, f"sheet_{s + 1:02d}.png")
            sheet.get_pixmap(matrix=pymupdf.Matrix(0.45, 0.45)).save(path)
            print(path)
    print(f"pages={d.page_count}")


def cmd_build(cfg, out, chunks, pdf=False):
    tmp = out + ".tmp.docx"
    with contextlib.redirect_stdout(io.StringIO()):
        cmd_init(cfg, tmp)
        for ch in chunks:
            cmd_add(tmp, ch)
    try:
        os.replace(tmp, out)
    except PermissionError:
        os.remove(tmp)
        sys.exit(f"error: {out} is locked (open in Word?). Close it and rerun build.")
    cmd_fix(out, pdf)


if __name__ == "__main__":
    cmd, *rest = sys.argv[1:]
    flags = [x for x in rest if x.startswith("--")]
    rest = [x for x in rest if not x.startswith("--")]
    {"init": lambda: cmd_init(*rest[:2]), "add": lambda: cmd_add(*rest[:2]),
     "fix": lambda: cmd_fix(rest[0], "--pdf" in flags),
     "build": lambda: cmd_build(rest[0], rest[1], rest[2:], "--pdf" in flags),
     "preview": lambda: cmd_preview(*rest[:3])}[cmd]()
